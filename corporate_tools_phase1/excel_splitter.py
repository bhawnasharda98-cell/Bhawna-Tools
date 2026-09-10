"""Split an Excel workbook into separate worksheet or row-chunk files."""

from __future__ import annotations

import argparse
import re
from pathlib import Path


def safe_name(value: str) -> str:
    cleaned = re.sub(r"[^A-Za-z0-9._-]+", "_", value.strip())
    return cleaned.strip("_") or "sheet"


def copy_rows(target, rows) -> None:
    for target_row, source_row in enumerate(rows, 1):
        for target_column, cell in enumerate(source_row, 1):
            target.cell(row=target_row, column=target_column).value = cell.value


def split_excel(input_file: Path, output_dir: Path, rows_per_file: int | None = None) -> list[Path]:
    try:
        from openpyxl import load_workbook, Workbook
    except ImportError as exc:
        raise SystemExit("Missing dependency: pip install openpyxl") from exc

    workbook = load_workbook(input_file)
    output_dir.mkdir(parents=True, exist_ok=True)
    written: list[Path] = []

    chunk_size = max(1, int(rows_per_file)) if rows_per_file else None
    for sheet_name in workbook.sheetnames:
        source = workbook[sheet_name]
        rows = list(source.iter_rows())
        if not chunk_size or len(rows) <= chunk_size:
            new_workbook = Workbook()
            target = new_workbook.active
            target.title = sheet_name[:31]
            copy_rows(target, rows)
            output_file = output_dir / f"{safe_name(sheet_name)}.xlsx"
            new_workbook.save(output_file)
            written.append(output_file)
            continue

        header = rows[:1]
        data_rows = rows[1:] if header else rows
        for index in range(0, len(data_rows), chunk_size):
            new_workbook = Workbook()
            target = new_workbook.active
            target.title = sheet_name[:31]
            chunk_number = index // chunk_size + 1
            chunk_rows = header + data_rows[index : index + chunk_size]
            copy_rows(target, chunk_rows)
            output_file = output_dir / f"{safe_name(sheet_name)}_part_{chunk_number}.xlsx"
            new_workbook.save(output_file)
            written.append(output_file)

    return written


def main() -> None:
    parser = argparse.ArgumentParser(description="Split an Excel workbook into worksheet files or row chunks.")
    parser.add_argument("input", type=Path, help="Path to the .xlsx workbook")
    parser.add_argument("--output-dir", type=Path, default=Path("excel_split_output"))
    parser.add_argument("--rows-per-file", type=int, help="Maximum data rows per split file. Header row is repeated.")
    args = parser.parse_args()

    files = split_excel(args.input, args.output_dir, args.rows_per_file)
    for file in files:
        print(file)


if __name__ == "__main__":
    main()

